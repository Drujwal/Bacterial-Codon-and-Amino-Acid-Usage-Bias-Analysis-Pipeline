#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Single-run bacterial CDS codon/protein usage pipeline.
Default input: newest .fa/.fasta/.fna/.fas file in the same directory as this script
Default output: Codon_Usage_Analysis.xlsx
Workbook sheets: 1_Codon_Usage and 2_Protein_Indices only

IMPORTANT:
The final GENOME_MEAN row in BOTH worksheets is calculated strictly as the
unweighted arithmetic mean of the individual accepted gene/CDS values
shown in the final worksheet.

No pooled counts, genome-level recalculation, or weighted averaging is used.
"""

import os
import re
import sys
import math
import traceback
from pathlib import Path
from collections import Counter, defaultdict

try:
    import numpy as np
    import pandas as pd
    from Bio import SeqIO
    from Bio.Seq import Seq
    from Bio.Data import CodonTable
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    print("A required package is missing:", exc)
    print("Install requirements with:")
    print("py -m pip install biopython pandas numpy openpyxl")
    input("Press Enter to exit...")
    raise

INPUT_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = INPUT_DIR / "Codon_Usage_Analysis.xlsx"

MIN_LENGTH_BP = 300
GENETIC_CODE_ID = 11

RIBOSOMAL_PATTERN = re.compile(
    r"ribosomal|ribosome|30s|50s|rps[a-z0-9]+|rpl[a-z0-9]+|"
    r"ribosomal protein|small subunit protein|large subunit protein",
    re.IGNORECASE
)

table = CodonTable.unambiguous_dna_by_id[GENETIC_CODE_ID]
STOP_CODONS = set(table.stop_codons)
SENSE_CODONS = sorted(table.forward_table.keys())
BASES = "TCAG"
ALL_CODONS = [a+b+c for a in BASES for b in BASES for c in BASES]

AA_ORDER = list("ACDEFGHIKLMNPQRSTVWY")

AA3 = {
    "A":"Ala","C":"Cys","D":"Asp","E":"Glu","F":"Phe","G":"Gly",
    "H":"His","I":"Ile","K":"Lys","L":"Leu","M":"Met","N":"Asn",
    "P":"Pro","Q":"Gln","R":"Arg","S":"Ser","T":"Thr","V":"Val",
    "W":"Trp","Y":"Tyr","*":"Stop"
}

AA_GROUPS = defaultdict(list)

for codon, aa in table.forward_table.items():
    AA_GROUPS[aa].append(codon)

SYN_FAMILIES = {
    aa: sorted(cods)
    for aa, cods in AA_GROUPS.items()
    if len(cods) > 1
}

# Average amino-acid residue masses (Da), after loss of water during
# peptide-bond formation. One H2O is added for the complete chain.
RESIDUE_MASS = {
    "A":71.0788,"R":156.1875,"N":114.1038,"D":115.0886,"C":103.1388,
    "E":129.1155,"Q":128.1307,"G":57.0519,"H":137.1411,"I":113.1594,
    "L":113.1594,"K":128.1741,"M":131.1926,"F":147.1766,"P":97.1167,
    "S":87.0782,"T":101.1051,"W":186.2132,"Y":163.1760,"V":99.1326
}

KD_HYDRO = {
    "A":1.8,"R":-4.5,"N":-3.5,"D":-3.5,"C":2.5,"Q":-3.5,"E":-3.5,
    "G":-0.4,"H":-3.2,"I":4.5,"L":3.8,"K":-3.9,"M":1.9,"F":2.8,
    "P":-1.6,"S":-0.8,"T":-0.7,"W":-0.9,"Y":-1.3,"V":4.2
}

PK = {
    "Cterm":3.55, "Nterm":7.50, "D":3.90, "E":4.07, "C":8.18,
    "Y":10.46, "H":6.04, "K":10.54, "R":12.48
}


def parse_header(record):
    desc = record.description or record.id

    def grab(pattern):
        m = re.search(pattern, desc, re.IGNORECASE)
        return m.group(1) if m else ""

    gene = grab(r"\bgene:([^\s]+)")
    symbol = grab(r"\bgene_symbol:([^\s]+)")
    protein_id = record.id
    product = grab(r"\bdescription:(.*)$")

    if not product:
        product = desc

    return desc, gene, symbol, product, protein_id


def calc_pi(seq):
    """Estimate pI by charge balance using a standard Henderson-Hasselbalch model."""
    if not seq:
        return np.nan

    counts = Counter(seq)

    def charge(ph):
        pos = counts.get("K",0)/(1+10**(ph-PK["K"]))
        pos += counts.get("R",0)/(1+10**(ph-PK["R"]))
        pos += counts.get("H",0)/(1+10**(ph-PK["H"]))
        pos += 1/(1+10**(ph-PK["Nterm"]))

        neg = counts.get("D",0)/(1+10**(PK["D"]-ph))
        neg += counts.get("E",0)/(1+10**(PK["E"]-ph))
        neg += counts.get("C",0)/(1+10**(PK["C"]-ph))
        neg += counts.get("Y",0)/(1+10**(PK["Y"]-ph))
        neg += 1/(1+10**(PK["Cterm"]-ph))

        return pos-neg

    lo, hi = 0.0, 14.0

    for _ in range(100):
        mid = (lo+hi)/2

        if charge(mid) > 0:
            lo = mid
        else:
            hi = mid

    return (lo+hi)/2


def protein_properties(protein):
    n = len(protein)

    if not n:
        return {}

    mw = sum(
        RESIDUE_MASS[aa]
        for aa in protein
    ) + 18.01528

    gravy = (
        sum(KD_HYDRO.get(aa, 0) for aa in protein) / n
    )

    aromo = (
        sum(protein.count(aa) for aa in "FWY") / n
    )

    return {
        "L_aa": n,
        "Molecular_Weight_KDa": mw/1000,
        "Gravy": gravy,
        "Aromo": aromo,
        "Isoelectric_Point": calc_pi(protein)
    }


def gc_fraction(seq):
    return (
        (seq.count("G")+seq.count("C"))/len(seq)
        if seq else np.nan
    )


def count_codons(seq):
    return Counter(
        seq[i:i+3]
        for i in range(0, len(seq)-2, 3)
    )


def calc_enc(codon_counts):
    """Wright-style ENC estimate from synonymous-family homozygosity."""

    fam_f = {
        2: [],
        3: [],
        4: [],
        6: []
    }

    for aa, codons in AA_GROUPS.items():

        k = len(codons)

        if k not in fam_f:
            continue

        n = sum(
            codon_counts.get(c,0)
            for c in codons
        )

        if n <= 1:
            continue

        freqs = [
            codon_counts.get(c,0)/n
            for c in codons
        ]

        F = (
            (sum(x*x for x in freqs) - 1/n)
            / (1 - 1/n)
        )

        F = min(
            1.0,
            max(1.0/k, F)
        )

        fam_f[k].append(F)

    def mean_f(k):
        return (
            float(np.mean(fam_f[k]))
            if fam_f[k]
            else np.nan
        )

    terms = [2.0]

    for k, coeff in [
        (2,9),
        (3,1),
        (4,5),
        (6,3)
    ]:

        f = mean_f(k)

        if np.isfinite(f) and f > 0:
            terms.append(coeff/f)

    return (
        float(sum(terms))
        if len(terms) > 1
        else np.nan
    )


def get_gc_positions(seq):
    """
    Calculate whole-CDS GC, GC12, and synonymous-site metrics.

    T3s/C3s/A3s/G3s/GC3s use only codons in synonymous amino-acid families.
    GC12 is the GC fraction across positions 1 and 2 of all complete codons.
    """

    codons = [
        seq[i:i+3]
        for i in range(0, len(seq)-2, 3)
    ]

    if not codons:
        return {}

    p1 = "".join(c[0] for c in codons)
    p2 = "".join(c[1] for c in codons)

    gc12 = (
        p1.count("G") + p1.count("C") +
        p2.count("G") + p2.count("C")
    ) / (2 * len(codons))

    silent_thirds = []

    for codon in codons:

        aa = table.forward_table.get(codon)

        if (
            aa is not None
            and len(AA_GROUPS[aa]) > 1
        ):
            silent_thirds.append(codon[2])

    n_silent = len(silent_thirds)

    if n_silent:

        t3 = silent_thirds.count("T") / n_silent
        c3 = silent_thirds.count("C") / n_silent
        a3 = silent_thirds.count("A") / n_silent
        g3 = silent_thirds.count("G") / n_silent
        gc3 = g3 + c3

        parity_at = (
            a3 / (a3 + t3)
            if a3 + t3
            else np.nan
        )

        parity_gc = (
            g3 / (g3 + c3)
            if g3 + c3
            else np.nan
        )

    else:

        t3 = c3 = a3 = g3 = gc3 = np.nan
        parity_at = parity_gc = np.nan

    return {
        "T3": t3,
        "C3": c3,
        "A3": a3,
        "G3": g3,
        "GC12": gc12,
        "GC3": gc3,
        "GC": gc_fraction(seq),
        "Parity_A3_over_A3plusT3": parity_at,
        "Parity_G3_over_G3plusC3": parity_gc,
        "Synonymous_Third_Sites": n_silent,
    }


def rscu_values(codon_counts):
    result = {}

    for aa, codons in AA_GROUPS.items():

        total = sum(
            codon_counts.get(c,0)
            for c in codons
        )

        k = len(codons)

        for c in codons:

            result[c] = (
                codon_counts.get(c,0)*k/total
                if total
                else np.nan
            )

    for c in STOP_CODONS:
        result[c] = np.nan

    return result


def codon_bias_metrics(
    codon_counts,
    cai_weights,
    preferred_codons
):
    """Calculate Fop, ENC and CAI."""

    rscu = rscu_values(codon_counts)

    synonymous_total = sum(
        codon_counts.get(c,0)
        for family in SYN_FAMILIES.values()
        for c in family
    )

    fop_num = sum(
        codon_counts.get(c,0)
        for c in preferred_codons
        if (
            c in SENSE_CODONS
            and any(
                c in family
                for family in SYN_FAMILIES.values()
            )
        )
    )

    fop = (
        fop_num / synonymous_total
        if synonymous_total and preferred_codons
        else np.nan
    )

    logsum = 0.0
    cai_n = 0
    cai_undefined = False
    cai_zero_weight = False

    for family in SYN_FAMILIES.values():

        observed_family_total = sum(
            codon_counts.get(c,0)
            for c in family
        )

        if not observed_family_total:
            continue

        if any(
            c not in cai_weights
            or not np.isfinite(cai_weights[c])
            for c in family
        ):
            cai_undefined = True
            break

        for c in family:

            n = codon_counts.get(c,0)

            if not n:
                continue

            weight = cai_weights[c]
            cai_n += n

            if weight <= 0:
                cai_zero_weight = True
            else:
                logsum += n * math.log(weight)

    if cai_undefined or not cai_n:

        cai = np.nan

    elif cai_zero_weight:

        cai = 0.0

    else:

        cai = math.exp(
            logsum / cai_n
        )

    return {
        "Fop": fop,
        "ENC": calc_enc(codon_counts),
        "CAI": cai
    }, rscu


def find_input():

    candidates = []

    for p in INPUT_DIR.iterdir():

        if (
            p.is_file()
            and p.suffix.lower()
            in {".fa",".fasta",".fna",".fas"}
        ):

            if p.name.lower() not in {
                "bacterial_codon_pipeline.fa"
            }:

                candidates.append(p)

    if not candidates:

        raise FileNotFoundError(
            "No .fa/.fasta/.fna/.fas file found in the script directory."
        )

    candidates.sort(
        key=lambda p: p.stat().st_mtime,
        reverse=True
    )

    if len(candidates) > 1:

        print(
            "Multiple FASTA files found. Using newest:",
            candidates[0].name
        )

        print(
            "To use another file, move other FASTA files out "
            "of this folder or edit find_input()."
        )

    return candidates[0]


def main():

    fasta = find_input()

    print(f"Input: {fasta}")

    records = list(
        SeqIO.parse(
            str(fasta),
            "fasta"
        )
    )

    if not records:

        raise ValueError(
            "No FASTA records found."
        )

    valid = []
    rejected = []
    stop_summary = Counter()

    for rec in records:

        desc, gene, symbol, product, pid = parse_header(rec)

        seq = (
            re.sub(
                r"\s+",
                "",
                str(rec.seq)
            )
            .upper()
            .replace("U", "T")
        )

        reason = ""

        if len(seq) < MIN_LENGTH_BP:

            reason = (
                f"Length < {MIN_LENGTH_BP} bp"
            )

        elif not seq or re.search(
            r"[^ACGT]",
            seq
        ):

            reason = (
                "Ambiguous/non-ACGT nucleotide"
            )

        elif len(seq) % 3 != 0:

            reason = (
                "Length not divisible by 3 (partial CDS)"
            )

        else:

            raw_codons = [
                seq[i:i+3]
                for i in range(0, len(seq), 3)
            ]

            terminal = (
                raw_codons[-1]
                if raw_codons
                and raw_codons[-1] in STOP_CODONS
                else ""
            )

            internal = (
                raw_codons[:-1]
                if terminal
                else raw_codons
            )

            if any(
                c in STOP_CODONS
                for c in internal
            ):

                reason = (
                    "Internal stop codon"
                )

        if reason:

            rejected.append({
                "Sequence_ID": rec.id,
                "Header": desc,
                "Length_bp": len(seq),
                "Reason": reason
            })

            continue

        raw_codons = [
            seq[i:i+3]
            for i in range(0, len(seq), 3)
        ]

        terminal_stop = (
            raw_codons[-1]
            if raw_codons
            and raw_codons[-1] in STOP_CODONS
            else ""
        )

        complete_seq = seq

        coding_seq = (
            complete_seq[:-3]
            if terminal_stop
            else complete_seq
        )

        codons = [
            coding_seq[i:i+3]
            for i in range(0, len(coding_seq), 3)
        ]

        protein = str(
            Seq(coding_seq).translate(
                table=GENETIC_CODE_ID,
                to_stop=False
            )
        )

        if "*" in protein:

            rejected.append({
                "Sequence_ID": rec.id,
                "Header": desc,
                "Length_bp": len(seq),
                "Reason": "Internal stop after translation"
            })

            continue

        stop_summary[
            terminal_stop or "No_terminal_stop"
        ] += 1

        valid.append({
            "record": rec,
            "description": desc,
            "gene": gene,
            "symbol": symbol,
            "product": product,
            "id": pid,
            "seq": seq,
            "coding_seq": coding_seq,
            "codons": codons,
            "codon_counts": Counter(codons),
            "protein": protein,
            "terminal_stop": terminal_stop,
            "is_ribosomal": bool(
                RIBOSOMAL_PATTERN.search(desc)
            )
        })

    if not valid:

        raise ValueError(
            "No sequences passed the filters. "
            "Check the input and filter rules."
        )

    # ------------------------------------------------------------
    # Ribosomal reference for CAI/Fop
    # ------------------------------------------------------------

    ribo_counts = Counter()

    ribo_n = sum(
        1
        for item in valid
        if item["is_ribosomal"]
    )

    for item in valid:

        if item["is_ribosomal"]:

            ribo_counts.update(
                item["codon_counts"]
            )

    cai_weights = {}
    preferred_codons = set()

    if ribo_n:

        for aa, family in SYN_FAMILIES.items():

            vals = {
                c: ribo_counts.get(c, 0)
                for c in family
            }

            maxv = max(
                vals.values()
            ) if vals else 0

            if maxv > 0:

                for c in family:

                    cai_weights[c] = (
                        vals[c] / maxv
                    )

                preferred_codons.update(
                    c
                    for c in family
                    if vals[c] == maxv
                )

    else:

        print(
            "WARNING: no ribosomal genes recognized from FASTA headers; "
            "CAI is blank and Fop is not reference-based."
        )

    codon_rows = []
    protein_rows = []
    rscu_matrix = []

    for item in valid:

        coding_seq = item["coding_seq"]
        codons = item["codons"]
        counts = item["codon_counts"]
        protein = item["protein"]

        gc = get_gc_positions(
            coding_seq
        )

        metrics, rscu = codon_bias_metrics(
            counts,
            cai_weights,
            preferred_codons
        )

        # P2 = (WWC + SST)/(WWY + SSY)
        wwc = 0
        sst = 0
        wwy = 0
        ssy = 0

        for codon in codons:

            first_two = codon[:2]
            third = codon[2]

            if all(
                base in "AT"
                for base in first_two
            ):

                if third in "CT":
                    wwy += 1

                if third == "C":
                    wwc += 1

            if all(
                base in "GC"
                for base in first_two
            ):

                if third in "CT":
                    ssy += 1

                if third == "T":
                    sst += 1

        p2_den = wwy + ssy

        p2 = (
            (wwc + sst) / p2_den
            if p2_den
            else np.nan
        )

        row = {
            "Sequence_ID": item["id"],
            "Gene": item["gene"],
            "Gene_Symbol": item["symbol"],
            "Product": item["product"],
            "Length_bp_including_terminal_stop": len(item["seq"]),
            "Ignored_3prime_incomplete_nt": 0,
            "Analyzed_CDS_bp_excluding_terminal_stop": len(coding_seq),
            "Codons_Analyzed": len(codons),
            "Terminal_Stop_Codon": (
                item["terminal_stop"]
                or "Absent"
            ),
            "Ribosomal_Reference": item["is_ribosomal"],
            **gc,
            **metrics,
            "P2": p2,
            "P2_Numerator_WWC_plus_SST": wwc + sst,
            "P2_Denominator_WWY_plus_SSY": p2_den,
            "Terminal_Stop_TAA": int(
                item["terminal_stop"] == "TAA"
            ),
            "Terminal_Stop_TAG": int(
                item["terminal_stop"] == "TAG"
            ),
            "Terminal_Stop_TGA": int(
                item["terminal_stop"] == "TGA"
            )
        }

        total_codons = len(codons)

        for codon in SENSE_CODONS:

            row[f"Count_{codon}"] = counts.get(
                codon,
                0
            )

        for codon in SENSE_CODONS:

            row[f"Per1000_{codon}"] = (
                1000 * counts.get(codon, 0)
                / total_codons
                if total_codons
                else np.nan
            )

        for codon in SENSE_CODONS:

            row[f"RSCU_{codon}"] = rscu.get(
                codon,
                np.nan
            )

        codon_rows.append(row)

        rscu_matrix.append({
            codon: rscu.get(
                codon,
                np.nan
            )
            for codon in SENSE_CODONS
        })

        props = protein_properties(
            protein
        )

        # L_sym = number of codons encoding amino acids
        # that have synonymous alternatives.
        l_sym = sum(
            1
            for codon in codons
            if (
                codon in table.forward_table
                and len(
                    AA_GROUPS[
                        table.forward_table[codon]
                    ]
                ) > 1
            )
        )

        protein_row = {
            "Sequence_ID": item["id"],
            "Gene": item["gene"],
            "Gene_Symbol": item["symbol"],
            "Product": item["product"],
            "Analyzed_CDS_Length_bp": len(coding_seq),
            "L_sym": l_sym,
            **props
        }

        protein_rows.append(
            protein_row
        )

    # ------------------------------------------------------------
    # PCA axes
    # ------------------------------------------------------------

    rscu_df = (
        pd.DataFrame(
            rscu_matrix,
            columns=SENSE_CODONS
        )
        .replace(
            [np.inf, -np.inf],
            np.nan
        )
    )

    rscu_df = (
        rscu_df
        .fillna(
            rscu_df.mean()
        )
        .fillna(0)
    )

    def pca_two_axes(df):

        if len(df) < 2 or df.shape[1] < 2:

            return np.full(
                (len(df), 2),
                np.nan
            )

        x = df.to_numpy(
            dtype=float
        )

        x = x - np.mean(
            x,
            axis=0
        )

        sd = np.std(
            x,
            axis=0
        )

        sd[sd == 0] = 1

        x = x / sd

        try:

            u, singular_values, _ = np.linalg.svd(
                x,
                full_matrices=False
            )

            n_axes = min(
                2,
                u.shape[1],
                len(singular_values)
            )

            scores = (
                u[:, :n_axes]
                * singular_values[:n_axes]
            )

            if scores.shape[1] == 1:

                scores = np.column_stack([
                    scores[:, 0],
                    np.zeros(len(scores))
                ])

            return scores[:, :2]

        except Exception:

            return np.full(
                (len(df), 2),
                np.nan
            )

    rscu_axes = pca_two_axes(
        rscu_df
    )

    aa_matrix = []

    for item in valid:

        aa_counts = Counter(
            item["protein"]
        )

        length = len(
            item["protein"]
        )

        aa_matrix.append([
            aa_counts.get(aa, 0) / length
            if length
            else 0
            for aa in AA_ORDER
        ])

    raau_axes = pca_two_axes(
        pd.DataFrame(
            aa_matrix,
            columns=AA_ORDER
        )
    )

    for i, row in enumerate(
        codon_rows
    ):

        row["RSCU_Axis1"] = (
            rscu_axes[i, 0]
        )

        row["RSCU_Axis2"] = (
            rscu_axes[i, 1]
        )

        row["RAAU_Axis1"] = (
            raau_axes[i, 0]
        )

        row["RAAU_Axis2"] = (
            raau_axes[i, 1]
        )

        protein_rows[i]["RAAU_Axis1"] = (
            raau_axes[i, 0]
        )

        protein_rows[i]["RAAU_Axis2"] = (
            raau_axes[i, 1]
        )

    codon_df = pd.DataFrame(
        codon_rows
    )

    protein_df = pd.DataFrame(
        protein_rows
    )

    # ------------------------------------------------------------
    # Final displayed columns
    # ------------------------------------------------------------

    codon_df["Genes"] = (
        codon_df["Gene_Symbol"]
        .replace("", np.nan)
        .fillna(
            codon_df["Gene"].replace(
                "",
                np.nan
            )
        )
        .fillna(
            codon_df["Sequence_ID"]
        )
    )

    codon_columns = {
        "Genes": "Genes",
        "T3": "T3s",
        "C3": "C3s",
        "A3": "A3s",
        "G3": "G3s",
        "ENC": "Nc",
        "GC3": "GC3s",
        "GC": "GC",
        "RSCU_Axis1": "Axis1(RSCU)",
        "RSCU_Axis2": "Axis2(RSCU)",
        "GC12": "GC12",
        "Parity_A3_over_A3plusT3":
            "Parity_A3/(A3+T3)",
        "Parity_G3_over_G3plusC3":
            "Parity_G3/(G3+C3)",
        "P2": "p2",
        "CAI": "CAI",
        "Fop": "Fop"
    }

    codon_df = (
        codon_df[
            list(codon_columns.keys())
        ]
        .rename(
            columns=codon_columns
        )
    )

    protein_df["Genes"] = (
        protein_df["Gene_Symbol"]
        .replace("", np.nan)
        .fillna(
            protein_df["Gene"].replace(
                "",
                np.nan
            )
        )
        .fillna(
            protein_df["Sequence_ID"]
        )
    )

    protein_columns = {
        "Genes": "Genes",
        "Isoelectric_Point":
            "Isoelectric_Point",
        "Molecular_Weight_KDa":
            "Molecular_Weight_KDa",
        "RAAU_Axis1":
            "Axis1(RAAU)",
        "RAAU_Axis2":
            "Axis2(RAAU)",
        "Gravy": "Gravy",
        "Aromo": "Aromo",
        "L_sym": "L_sym",
        "L_aa": "L_aa"
    }

    protein_df = (
        protein_df[
            list(protein_columns.keys())
        ]
        .rename(
            columns=protein_columns
        )
    )

    # ============================================================
    # FINAL ARITHMETIC MEAN
    # ============================================================
    #
    # IMPORTANT:
    # The summary is created AFTER the final displayed columns
    # have been selected and renamed.
    #
    # Therefore every numeric value in GENOME_MEAN is calculated
    # directly from the exact individual-gene values that appear
    # above it in the final Excel worksheet.
    #
    # This is mathematically equivalent to:
    #
    #     =AVERAGE(B2:B3500)
    #
    # when rows 2:3500 are the individual accepted genes.
    #
    # No pooled calculation.
    # No genome-level recalculation.
    # No weighted average.
    # ============================================================

    def add_arithmetic_mean_row(df):

        summary = {}

        for col in df.columns:

            if pd.api.types.is_numeric_dtype(
                df[col]
            ):

                summary[col] = df[col].mean(
                    skipna=True
                )

            else:

                summary[col] = ""

        summary["Genes"] = "GENOME_MEAN"

        return pd.concat(
            [
                df,
                pd.DataFrame([summary])
            ],
            ignore_index=True
        )

    codon_gene_count = len(codon_df)
    protein_gene_count = len(protein_df)

    if codon_gene_count != protein_gene_count:

        raise RuntimeError(
            "Codon and protein gene counts do not match."
        )

    codon_df = add_arithmetic_mean_row(
        codon_df
    )

    protein_df = add_arithmetic_mean_row(
        protein_df
    )

    # ------------------------------------------------------------
    # Final row-count validation
    # ------------------------------------------------------------

    expected_rows = len(valid) + 1

    for name, frame in [
        ("1_Codon_Usage", codon_df),
        ("2_Protein_Indices", protein_df)
    ]:

        if len(frame) != expected_rows:

            raise RuntimeError(
                f"Internal row-count check failed for "
                f"{name}: {len(frame)} != {expected_rows}"
            )

    # ------------------------------------------------------------
    # Excel output
    # ------------------------------------------------------------

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl"
    ) as writer:

        codon_df.to_excel(
            writer,
            sheet_name="1_Codon_Usage",
            index=False
        )

        protein_df.to_excel(
            writer,
            sheet_name="2_Protein_Indices",
            index=False
        )

        wb = writer.book

        for ws in wb.worksheets:

            ws.freeze_panes = "A2"

            ws.auto_filter.ref = (
                ws.dimensions
                if ws.max_row > 1
                else "A1"
            )

            # Header formatting
            for cell in ws[1]:

                cell.font = Font(
                    bold=True,
                    color="FFFFFF"
                )

                cell.fill = PatternFill(
                    "solid",
                    fgColor="1F4E78"
                )

                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True
                )

            # Column widths
            for col_cells in ws.columns:

                col_letter = get_column_letter(
                    col_cells[0].column
                )

                max_len = 0

                for cell in col_cells[
                    :min(ws.max_row, 300)
                ]:

                    value = cell.value

                    if value is not None:

                        max_len = max(
                            max_len,
                            len(str(value))
                        )

                ws.column_dimensions[
                    col_letter
                ].width = min(
                    max(max_len + 2, 12),
                    32
                )

            # Final row = arithmetic mean summary
            summary_row = ws.max_row

            for cell in ws[summary_row]:

                cell.font = Font(
                    bold=True,
                    color="000000"
                )

                cell.fill = PatternFill(
                    "solid",
                    fgColor="D9EAF7"
                )

                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True
                )

            ws.sheet_view.showGridLines = True

    # ------------------------------------------------------------
    # Rejected-record log
    # ------------------------------------------------------------

    if rejected:

        rejected_path = (
            INPUT_DIR /
            "Codon_Usage_Analysis_rejected_records.csv"
        )

        pd.DataFrame(
            rejected
        ).to_csv(
            rejected_path,
            index=False
        )

        print(
            "Rejected-record log:",
            rejected_path
        )

    # ------------------------------------------------------------
    # Final messages
    # ------------------------------------------------------------

    print("\nANALYSIS COMPLETE")

    print(
        "Input records:",
        len(records)
    )

    print(
        "Analyzed CDS/protein records:",
        len(valid)
    )

    print(
        "Rows in each worksheet including "
        "one arithmetic-mean summary row:",
        len(valid) + 1
    )

    print(
        "Rejected records:",
        len(rejected)
    )

    print(
        "Ribosomal reference records detected:",
        ribo_n
    )

    print(
        "\nFINAL SUMMARY ROW:"
        "\n  Every numeric value is the unweighted arithmetic mean"
        "\n  of the individual accepted gene/CDS values shown"
        "\n  in the final worksheet."
        "\n  No pooled or genome-level recalculation is used."
    )

    print(
        "\nWorkbook contains exactly two sheets:"
        "\n  1_Codon_Usage"
        "\n  2_Protein_Indices"
    )

    print(
        "\nExcel output:",
        OUTPUT_FILE
    )

    if not ribo_n:

        print(
            "\nWARNING: no ribosomal genes detected; "
            "CAI is blank and Fop has no reference-defined "
            "preferred codons."
        )

    print(
        "\nNote: Nc is a Wright-style estimate with "
        "finite-sample correction; verify selected records "
        "against a trusted implementation before publication."
    )

    input("\nPress Enter to close...")


if __name__ == "__main__":

    try:

        main()

    except Exception:

        print("\nPIPELINE FAILED:")

        traceback.print_exc()

        input("\nPress Enter to close...")